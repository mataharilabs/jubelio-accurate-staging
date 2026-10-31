"""Rapikan Laporan Penjualan Jubelio -> format staging (v2).
Usage: python3 rapikan.py <raw.xlsx> <bobot.xlsx> <out.xlsx> [--rebuild-mapping]
Mapping provinsi/kota dibaca dari mapping_provinsi.xlsx / mapping_kota.xlsx (di folder yang sama
dengan script) bila ada: baris yang sudah ada (hasil review user) dipakai apa adanya, nilai baru
dicocokkan otomatis lalu ditambahkan. --rebuild-mapping = abaikan file lama.
"""
import pandas as pd, numpy as np, warnings, sys, os, re, json, difflib
warnings.filterwarnings('ignore')
argv=sys.argv[1:]; TANGGAL=None
if '--tanggal' in argv:   # --tanggal YYYY-MM-DD : hanya baris dengan Tanggal di hari itu (00:00:00-23:59:59 WIB)
    i=argv.index('--tanggal'); TANGGAL=pd.Timestamp(argv[i+1]).normalize(); del argv[i:i+2]
args=[a for a in argv if not a.startswith('--')]
src, bobotf, fn = args[:3]
REBUILD='--rebuild-mapping' in sys.argv
HERE=os.path.dirname(os.path.abspath(__file__))
MAP_PROV=os.path.join(HERE,'mapping_provinsi.xlsx'); MAP_KOTA=os.path.join(HERE,'mapping_kota.xlsx')
WIL=json.load(open(os.path.join(HERE,'wilayah','wilayah_kemendagri.json')))  # cahyadsn/wilayah, Kepmendagri 300.2.2-2138/2025

# ---------------- baris & aturan lama (tidak diubah) ----------------
df=pd.read_excel(src, dtype={'REF':str,'SKU':str,'item_code_bundle':str,'No Pesanan':str})
_bx=pd.ExcelFile(bobotf); b=pd.read_excel(bobotf, sheet_name=('Tabel Bundle' if 'Tabel Bundle' in _bx.sheet_names else 0), header=None, dtype=str).iloc[:,1:3].dropna()
b.columns=['key','bobot']; b=b[b['key']!='SKU BUNDLE']; b['key']=b['key'].str.strip(); b['bobot']=b['bobot'].astype(float)
bmap=dict(zip(b['key'],b['bobot']))
# --- mapping Channel / Store dari sheet 'Channel' & 'Store' di file bobot (match case/spasi-insensitive) ---
_nk=lambda v: re.sub(r'\s+',' ',str(v)).strip().lower()
def _ldmap(sheet):
    if sheet not in _bx.sheet_names: print('WARNING: sheet',sheet,'tidak ada -> nilai raw dipakai'); return {}
    m=pd.read_excel(bobotf,sheet_name=sheet,header=None,dtype=str).iloc[:,1:3].dropna()
    m=m.iloc[1:] if len(m) else m   # buang header
    m=m.drop_duplicates(); d={}
    for k,v in zip(m.iloc[:,0],m.iloc[:,1]):
        kk=_nk(k); v=str(v).strip()
        if kk in d and d[kk]!=v: print('WARNING: mapping',sheet,'konflik untuk',k,':',d[kk],'vs',v,'-> pakai yang pertama'); continue
        d.setdefault(kk,v)
    return d
CHMAP,STMAP=_ldmap('Channel'),_ldmap('Store')   # (lama, tidak dipakai lagi untuk output; diganti mapping_channel_toko_jubelio.xlsx)
# --- mapping Channel+Toko dari user (31/10): key = Channel raw + Nama Toko raw -> Channel & Nama Toko 'Website / Accurate' ---
MAP_CT=os.path.join(HERE,'mapping_channel_toko_jubelio.xlsx')
_ct=pd.read_excel(MAP_CT,header=None,dtype=str)
_hr=next(i for i in range(len(_ct)) if _ct.iloc[i].fillna('').str.strip().str.lower().tolist().count('channel')>=2)  # baris header 'Channel | Nama Toko | Channel | Nama Toko'
_cc=[j for j,v in enumerate(_ct.iloc[_hr]) if str(v).strip().lower() in ('channel','nama toko')][:4]
CTMAP={}
for r in _ct.iloc[_hr+1:,_cc].itertuples(index=False):
    if pd.isna(r[0]) or not str(r[0]).strip(): continue
    k=(_nk(r[0]), _nk(r[1]) if pd.notna(r[1]) else '')
    v=(str(r[2]).strip(), str(r[3]).strip())
    if k in CTMAP and CTMAP[k]!=v: print('WARNING: mapping channel/toko konflik',k,CTMAP[k],'vs',v,'-> pakai yang pertama'); continue
    CTMAP.setdefault(k,v)
print('mapping channel/toko:',MAP_CT,len(CTMAP),'pasangan')
def _apply(col,mp,label):
    out=col.map(lambda v: mp.get(_nk(v),v) if pd.notna(v) else v)
    un=sorted({str(v) for v in col.dropna() if _nk(v) not in mp})
    if un: print('WARNING: %s tidak ter-mapping (raw dipakai):'%label,un)
    vc=pd.DataFrame({'raw':col,'mapped':out}).value_counts(dropna=False).reset_index(name='rows')
    print('%s mapping applied (raw rows setelah filter status):'%label); print(vc.to_string(index=False))
    return out
df['_t']=pd.to_datetime(df['Tanggal'],format='mixed')
print('raw export Tanggal range:',df['_t'].min(),'->',df['_t'].max(),'rows',len(df))
if TANGGAL is not None:
    d0,d1=TANGGAL,TANGGAL+pd.Timedelta(days=1)
    full=df['_t'].min()<d0 and df['_t'].max()>=d1-pd.Timedelta(seconds=1)
    print('filter',d0.date(),'| hari tercakup penuh oleh export:',full)
    if not full: print('WARNING: export tidak mencakup seluruh hari',d0.date())
    df=df[(df['_t']>=d0)&(df['_t']<d1)].copy()
    print('filtered Tanggal range:',df['_t'].min(),'->',df['_t'].max())
n_raw=len(df)
df=df[~df['Status'].str.upper().isin(['PENDING','PAID','CANCELED','CANCELLED','CANCEL'])].copy()
for c in ['SKU','item_code_bundle','REF','No Pesanan']:   # buang karakter tak terlihat (mis. U+2060)
    df[c]=df[c].str.replace('[\u200b-\u200f\u2060\ufeff\xa0]','',regex=True)
for c in ['Qty','Harga Jual','Total Penjualan Per Produk','Diskon','Diskon Lainnya','Biaya Lainnya','qty_bundle']:
    df[c]=pd.to_numeric(df[c],errors='coerce').fillna(0)
isb=df['item_code_bundle'].notna() & (df['item_code_bundle'].str.strip()!='')
shp=df['Channel'].str.upper().str.strip().eq('SHOPEE')
ref=df['REF'].where(df['REF'].notna() & (df['REF'].str.strip()!=''), df['No Pesanan'])
key=(df['SKU'].str.strip()+'-'+df['item_code_bundle'].fillna('').str.strip()).where(isb,'')
w=key.map(bmap).where(isb,1.0)
# Unit Price per item (per pcs setelah bundle dipecah). Harga Jual = harga per 1 unit/bundle, jadi TIDAK dikali Qty raw
# (revisi 31/10: JM-20261026-V9QU7Q Qty 4 x 28.000 -> Unit Price 28.000). Untuk Qty raw = 1 hasilnya sama dengan aturan lama.
gross=(df['Total Penjualan Per Produk']*w/df['Qty']).where(shp, df['Harga Jual']*w)
order_disc=df['Diskon Lainnya'].where(shp, df['Diskon']+df['Diskon Lainnya']+df['Biaya Lainnya'])  # diskon level pesanan
miss=sorted(set(key[isb & w.isna()]))

# ---------------- mapping wilayah ----------------
def sq(s): return re.sub(r'[^a-z]','',str(s).lower())
PROV={k:n.strip() for k,n in WIL['prov']}
PROV['31']='DKI Jakarta'  # permintaan user: pakai 'DKI Jakarta'
KAB={k:n.strip() for k,n in WIL['kab']}
KEC=[(k,n.strip()) for k,n in WIL['kec']]
PFX=['kabupatenadministrasi','kotaadministrasi','kabupaten','kota']
def core(n):
    s=sq(n)
    for p in PFX:
        if s.startswith(p): return s[len(p):]
    return s
PALIAS={'dkijakarta':'31','jakarta':'31','dki':'31','daerahkhususjakarta':'31','diyogyakarta':'34','yogyakarta':'34',
 'daerahistimewayogyakarta':'34','jogja':'34','nanggroeacehdarussalam':'11','nad':'11','diaceh':'11','aceh':'11',
 'bangkabelitung':'19','babel':'19','ntb':'52','ntt':'53','eastjava':'35','westjava':'32','centraljava':'33',
 'jakartaraya':'31','sumut':'12','sumbar':'13','sumsel':'16','kalbar':'61','kalteng':'62','kalsel':'63','kaltim':'64',
 'kaltara':'65','sulut':'71','sulteng':'72','sulsel':'73','sultra':'74','sulbar':'76','kepri':'21'}
PSQ={sq(n):k for k,n in PROV.items()}
def match_prov(raw):
    if pd.isna(raw) or not str(raw).strip(): return None,'kosong',0.0
    s=sq(re.sub(r'\(.*?\)','',str(raw)))
    if s in PSQ: return PSQ[s],'exact',1.0
    if s in PALIAS: return PALIAS[s],'alias',1.0
    for kw,code in [('yogyakarta','34'),('jakarta','31'),('aceh','11'),('bangkabelitung','19')]:
        if kw in s: return code,'alias (kata kunci)',0.95
    best=max(PSQ,key=lambda p:difflib.SequenceMatcher(None,s,p).ratio()); sc=difflib.SequenceMatcher(None,s,best).ratio()
    return PSQ[best],'fuzzy',round(sc,3)

KMANUAL={'jakarta':'Jakarta','lombok':'Lombok'}  # keputusan user 29/10
KALIAS={'ldenpasar':'denpasar','kajenpekalongan':'kajen'}
# ibu kota kabupaten yang bukan nama kecamatan resmi
KCODE={'caruban':'35.19','meulaboh':'11.05','solo':'33.72'}
def match_kota(raw,pk):
    """return (kode kab/kota, metode, skor, catatan)"""
    if pd.isna(raw) or not str(raw).strip(): return 'MANUAL:Tanpa nama kota','aturan user',1.0,'Kota kosong -> "Tanpa nama kota"'
    if sq(raw) in KMANUAL: return 'MANUAL:'+KMANUAL[sq(raw)],'aturan user',1.0,'ditetapkan user'
    r=str(raw).strip(); s=sq(r)
    s=re.sub(r'^kab(?!upaten)','kabupaten',s)
    s=KALIAS.get(s,s)
    if s in KCODE and (pk is None or KCODE[s].startswith(pk+'.')): return KCODE[s],'alias manual',0.8,f'"{r}" = {KAB[KCODE[s]]}'
    kabs={k:n for k,n in KAB.items() if pk is None or k.startswith(pk+'.')}
    full={sq(n):k for k,n in kabs.items()}
    if s in full: return full[s],'exact',1.0,''
    # ditulis dengan prefix -> cocokkan tipe + nama inti
    for p,typ in [('kotaadministrasi','Kota'),('kabupatenadministrasi','Kabupaten'),('kabupaten','Kabupaten'),('kota','Kota')]:
        if s.startswith(p):
            c=[k for k,n in kabs.items() if core(n)==s[len(p):] and n.startswith(typ)]
            if len(c)==1: return c[0],'prefix+nama',1.0,''
    # tanpa prefix -> nama inti
    c=[k for k,n in kabs.items() if core(n)==s]
    if len(c)==1: return c[0],'nama inti',1.0,''
    if len(c)>1:
        kota=[k for k in c if KAB[k].startswith('Kota')]
        pick=kota[0] if kota else c[0]
        return pick,'ambigu kota/kab',0.5,'Ada '+' & '.join(KAB[k] for k in c)+'; dipilih Kota (tidak ada prefix di data)'
    # nama kecamatan (mis. Cibinong, Kepanjen) -> induk kab/kota
    kc=[k for k,n in KEC if (pk is None or k.startswith(pk+'.')) and sq(n)==s]
    parents=sorted(set(k[:5] for k in kc))
    if len(parents)==1: return parents[0],'kecamatan',0.8,f'"{r}" adalah kecamatan di {KAB[parents[0]]}'
    if len(parents)>1:
        return parents[0],'kecamatan ambigu',0.4,'Kecamatan ada di: '+', '.join(KAB[p] for p in parents)
    # awalan beberapa kab/kota (mis. JAKARTA -> 5 kota adm.) -> tidak bisa ditentukan
    pre=[k for k,n in kabs.items() if core(n).startswith(s)]
    if len(pre)>1: return None,'ambigu',0.0,'Cocok ke: '+', '.join(KAB[k] for k in pre)+'; perlu diisi manual'
    # awalan nama kecamatan (mis. CIKARANG -> Cikarang Utara/Selatan/... di Kab. Bekasi)
    kp=sorted(set(k[:5] for k,n in KEC if (pk is None or k.startswith(pk+'.')) and sq(n).startswith(s)))
    if len(kp)==1: return kp[0],'awalan kecamatan',0.8,f'kecamatan "{r} ..." ada di {KAB[kp[0]]}'
    # fuzzy atas nama kab/kota inti dan kecamatan
    cand=[(core(n),k,'fuzzy kab/kota') for k,n in kabs.items()]+[(sq(n),k[:5],'fuzzy kecamatan') for k,n in KEC if pk is None or k.startswith(pk+'.')]
    best=max(cand,key=lambda x:difflib.SequenceMatcher(None,s,x[0]).ratio())
    sc=round(difflib.SequenceMatcher(None,s,best[0]).ratio(),3)
    if sc<0.75: return None,'tidak cocok',sc,'Tidak ditemukan; perlu diisi manual'
    return best[1],best[2],sc,f'mirip "{best[0]}"'

def load_map(path,keys):
    if REBUILD or not os.path.exists(path): return None
    m=pd.read_excel(path,dtype=str).fillna('')
    return {tuple(row[k] for k in keys):row for _,row in m.iterrows()}

# provinsi
oldp=load_map(MAP_PROV,['Provinsi (Jubelio)'])
prows=[]
for raw in sorted(df['Provinsi'].fillna('').astype(str).unique()):
    if oldp and (raw,) in oldp: prows.append(dict(oldp[(raw,)])); continue
    k,m,sc=match_prov(raw)
    prows.append({'Provinsi (Jubelio)':raw,'Provinsi Resmi':PROV.get(k,''),'Kode':k or '','Metode':m,'Skor':sc,
                  'Perlu Review':'YA' if (sc<0.9 or not k) else ''})
pm=pd.DataFrame(prows)
if oldp:  # sertakan juga baris lama yang tidak muncul hari ini
    pm=pd.concat([pm,pd.DataFrame([dict(v) for kk,v in oldp.items() if kk[0] not in set(pm['Provinsi (Jubelio)'])])],ignore_index=True)
p2off=dict(zip(pm['Provinsi (Jubelio)'],pm['Provinsi Resmi'])); p2code=dict(zip(pm['Provinsi (Jubelio)'],pm['Kode']))
# kota
oldk=load_map(MAP_KOTA,['Provinsi (Jubelio)','Kota (Jubelio)'])
krows=[]
pairs=df[['Provinsi','Kota']].fillna('').astype(str).drop_duplicates().sort_values(['Provinsi','Kota'])
for pr,raw in pairs.itertuples(index=False):
    if oldk and (pr,raw) in oldk: krows.append(dict(oldk[(pr,raw)])); continue
    pk=p2code.get(pr) or None
    k,m,sc,note=match_kota(raw,pk)
    if k is None and pk:   # coba lintas provinsi
        k2,m2,sc2,n2=match_kota(raw,None)
        if k2 and sc2>=0.8: k,m,sc,note=k2,m2+' (lintas provinsi)',min(sc2,0.6),n2+'; provinsi di data berbeda'
    krows.append({'Provinsi (Jubelio)':pr,'Kota (Jubelio)':raw,'Provinsi Resmi':p2off.get(pr,''),'Kab/Kota Resmi':(k[7:] if k and k.startswith('MANUAL:') else KAB.get(k,'')),
                  'Kode':'' if (not k or k.startswith('MANUAL:')) else k,'Metode':m,'Skor':sc,'Perlu Review':'YA' if (sc<0.9 or not k) else '','Catatan':note})
km=pd.DataFrame(krows)
if oldk:
    seen=set(zip(km['Provinsi (Jubelio)'],km['Kota (Jubelio)']))
    km=pd.concat([km,pd.DataFrame([dict(v) for kk,v in oldk.items() if kk not in seen])],ignore_index=True)
k2off={(a,c):o for a,c,o in zip(km['Provinsi (Jubelio)'],km['Kota (Jubelio)'],km['Kab/Kota Resmi'])}

# ---------------- susun baris ----------------
def cust_code(n):
    l=re.sub(r'[^A-Za-z]','',str(n) if pd.notna(n) else '').upper()
    return (l+'XXX')[:3]
# Jubelio-POS = penjualan offline: kota dari Lokasi Gudang
POSCITY={'lintas sby':('Surabaya','Jawa Timur'),'lintas jkt':('Jakarta','DKI Jakarta')}
pos=df['Channel'].str.upper().str.strip().eq('JUBELIO-POS')
provraw=df['Provinsi'].fillna('').astype(str); kotaraw=df['Kota'].fillna('').astype(str)
# JUBELIO-POS: Nama Toko raw = nama lokasi (mis. 'Lintas Jkt') -> pakai baris file dengan toko kosong (Jubelio POS / Jansen Official Store)
_ckey=[(_nk(c) if pd.notna(c) else '', '' if p else (_nk(t) if pd.notna(t) else '')) for c,t,p in zip(df['Channel'],df['Nama Toko'],pos)]
_hit=[CTMAP.get(k) for k in _ckey]
ch_out=[h[0] if h else c for h,c in zip(_hit,df['Channel'])]
st_out=[h[1] if h else t for h,t in zip(_hit,df['Nama Toko'])]
UNMAPPED=pd.DataFrame({'No Pesanan':ref,'Channel raw':df['Channel'],'Nama Toko raw':df['Nama Toko'],'hit':[h is not None for h in _hit]})
UNMAPPED=UNMAPPED[~UNMAPPED.hit].groupby(['Channel raw','Nama Toko raw'],dropna=False)['No Pesanan'].agg(lambda x: sorted(set(x))).reset_index()
_vc=pd.DataFrame({'Channel raw':df['Channel'],'Toko raw':df['Nama Toko'].where(~pos,'(POS) '+df['Nama Toko'].fillna('')),'Channel':ch_out,'Store Name':st_out}).value_counts(dropna=False).reset_index(name='rows')
print('Channel/Store mapping applied (baris setelah filter status):'); print(_vc.to_string(index=False))
if len(UNMAPPED): print('WARNING: pasangan Channel/Toko TIDAK ter-mapping (nilai raw dipakai):'); print(UNMAPPED.to_string(index=False))
else: print('Semua pasangan Channel/Toko ter-mapping')
row=pd.DataFrame({'Tanggal Pesanan':df['_t'],'No Pesanan':ref,'Channel':ch_out,'Store Name':st_out,

 'SKU':df['item_code_bundle'].where(isb, df['SKU']).str.strip(),
 'Qty':(df['qty_bundle']*df['Qty']).where(isb, df['Qty']).astype(int),'Unit Price':gross,'_od':order_disc,
 'Customer Name':df['Pelanggan'],
 'Province':[p2off.get(p) or p for p in provraw],
 'City':[k2off.get((p,k)) or k for p,k in zip(provraw,kotaraw)],'Kurir':df['Kurir']})
for i in row.index[pos.values]:
    loc=str(df.at[i,'Lokasi Gudang']).strip().lower()
    hit=next((v for k,v in POSCITY.items() if loc.startswith(k)),None)   # termasuk varian 'Lintas Sby (Bazar)' dst.
    if hit:
        row.at[i,'City']=hit[0]
        if not str(df.at[i,'Provinsi'] if pd.notna(df.at[i,'Provinsi']) else '').strip(): row.at[i,'Province']=hit[1]
    else:
        print('WARNING: POS row lokasi tanpa aturan kota (tidak diubah):',df.at[i,'No Pesanan'],df.at[i,'Lokasi Gudang'])
print('POS rows overridden:',int(pos.sum()), df.loc[pos,['No Pesanan','Lokasi Gudang']].values.tolist())
row['_sub']=row['Qty']*row['Unit Price']
n_before=len(row)
# cek diskon level pesanan konsisten
dchk=row.groupby('No Pesanan')['_od'].nunique(); disc_diff=list(dchk[dchk>1].index)
# gabung SKU sama dalam 1 pesanan
first=['Tanggal Pesanan','Channel','Store Name','_od','Customer Name','Province','City','Kurir']
g=row.groupby(['No Pesanan','SKU'],sort=False)
out=g[first].first().join(g['Qty'].sum()).join(g['_sub'].apply(lambda x: x.sum() if x.notna().all() else np.nan)).join(g.size().rename('_n')).reset_index()   # bobot hilang -> Sub Total kosong, bukan 0
merged=out[out['_n']>1]
out['Unit Price']=np.where(out['Qty']!=0,out['_sub']/out['Qty'],0)
out['Sub Total']=out['_sub']
out['Total Nilai Pesanan']=out.groupby('No Pesanan')['Sub Total'].transform('sum')
out['Total Diskon']=out['_od']
out['Unit Diskon']=np.where((out['Total Nilai Pesanan']!=0)&(out['Qty']!=0),
    out['Sub Total']/out['Total Nilai Pesanan']*out['Total Diskon']/out['Qty'],0)
out['Revenue per Unit']=out['Unit Price']-out['Unit Diskon']   # pendapatan bersih per unit setelah diskon
out['Customer Code']=out['Customer Name'].map(cust_code)
COLS=['Tanggal Pesanan','No Pesanan','Channel','Store Name','SKU','Qty','Unit Price','Sub Total','Total Nilai Pesanan',
      'Total Diskon','Unit Diskon','Revenue per Unit','Customer Name','Customer Code','Province','City','Kurir']
final=out[COLS]

# ---------------- simpan ----------------
from openpyxl.styles import PatternFill, Font
def write(dfx,path,sheet,widths,numcols=(),datecol=None,flagcol=None):
    with pd.ExcelWriter(path,engine='openpyxl') as x:
        dfx.to_excel(x,index=False,sheet_name=sheet); ws=x.sheets[sheet]
        for i,wd in enumerate(widths): ws.column_dimensions[ws.cell(1,i+1).column_letter].width=wd
        for c in range(1,dfx.shape[1]+1): ws.cell(1,c).font=Font(bold=True)
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        yel=PatternFill('solid',fgColor='FFF2CC')
        for r in range(2,len(dfx)+2):
            for c in numcols: ws.cell(r,c).number_format='#,##0.##'
            if datecol: ws.cell(r,datecol).number_format='yyyy-mm-dd hh:mm'
            if flagcol and ws.cell(r,flagcol).value=='YA':
                for c in range(1,dfx.shape[1]+1): ws.cell(r,c).fill=yel
write(final,fn,'Data',[17,22,16,22,16,7,14,15,17,14,13,15,26,9,26,32,30],numcols=(7,8,9,10,11,12),datecol=1)
write(pm,MAP_PROV,'Provinsi',[34,34,8,10,8,12],flagcol=6)
write(km,MAP_KOTA,'Kota',[30,32,30,36,9,22,8,12,60],flagcol=8)

# ---------------- ringkasan & cek ----------------
chk=final.assign(x=final['Unit Diskon']*final['Qty'],y=final['Revenue per Unit']*final['Qty']).groupby('No Pesanan').agg(a=('x','sum'),r=('y','sum'),d=('Total Diskon','first'),t=('Total Nilai Pesanan','first'))
bad=chk[(chk.a-chk.d).abs()>0.01]
badr=chk[(chk.r-(chk.t-chk.d)).abs()>0.01]
print('raw',n_raw,'after status filter',n_before,'after merge',len(final),'merged groups',len(merged),'orders',final['No Pesanan'].nunique())
print('missing bobot',miss,'| orders w/ differing discount',disc_diff)
print('SUM Sub Total',round(final['Sub Total'].sum(),2),'| SUM Unit Price (pre-merge rows)',round(gross.sum(),2),'| SUM row Sub Total pre-merge',round(row['_sub'].sum(),2))
print('SUM Total Diskon (per order)',round(chk.d.sum(),2),'| SUM UnitDiskon*Qty',round(chk.a.sum(),2))
print('SUM Revenue per Unit*Qty',round(chk.r.sum(),2),'| SUM (Total Nilai Pesanan - Total Diskon)',round((chk.t-chk.d).sum(),2),'| orders where revenue check fails',len(badr)); print(badr.head(10).to_string())
print('orders where disc check fails',len(bad)); print(bad.head(10).to_string())
print('merged examples:'); print(merged.head(5).to_string())
print('prov review:'); print(pm[pm['Perlu Review']=='YA'].to_string())
print('kota review:'); print(km[km['Perlu Review']=='YA'][['Provinsi (Jubelio)','Kota (Jubelio)','Kab/Kota Resmi','Metode','Skor','Catatan']].to_string())
